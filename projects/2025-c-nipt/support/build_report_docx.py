"""Build the polished DOCX report from computed project artifacts.

Preset: standard_business_brief.
Named override: Chinese text uses Microsoft YaHei as the East Asian font;
the title is 26 pt #16324F with 0 pt before and 12 pt after.
"""

from __future__ import annotations

import csv
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
PAPER = ROOT / "paper"
FIGURES = ROOT / "figures"
TABLES = ROOT / "tables"
OUTPUT = PAPER / "report.docx"

PRESET = {
    "page_width": 8.2677,
    "page_height": 11.6929,
    "margin": 1.0,
    "header_footer": 0.492,
    "content_width_dxa": 9000,
    "table_indent_dxa": 120,
    "cell_margin_dxa": {"top": 80, "bottom": 80, "start": 120, "end": 120},
    "body_after_pt": 6,
    "body_line_spacing": 1.10,
    "h1": {"size": 16, "before": 16, "after": 8, "color": "2E74B5"},
    "h2": {"size": 13, "before": 12, "after": 6, "color": "2E74B5"},
    "h3": {"size": 12, "before": 8, "after": 4, "color": "1F4D78"},
    "header_fill": "F2F4F7",
    "callout_fill": "F4F6F9",
    "ink": "16324F",
    "muted": "5B6573",
    "border": "CBD3DC",
}


def read_csv(name: str) -> list[dict[str, str]]:
    with (TABLES / name).open("r", encoding="utf-8-sig", newline="") as handle:
        return list(csv.DictReader(handle))


def set_font(run, *, latin="Calibri", east_asia="Microsoft YaHei", size=None, bold=None, color=None):
    run.font.name = latin
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), east_asia)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)


def set_style_font(style, *, latin="Calibri", east_asia="Microsoft YaHei"):
    style.font.name = latin
    style._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), east_asia)


def set_cell_shading(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_width(cell, width_dxa: int):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_dxa))
    tc_w.set(qn("w:type"), "dxa")


def set_table_borders(table, color: str = "CBD3DC", size: str = "6"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = borders.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def set_table_geometry(table, widths_dxa: list[int], *, indent_dxa: int = 120):
    if sum(widths_dxa) != PRESET["content_width_dxa"]:
        raise ValueError(
            f"Table widths must sum to {PRESET['content_width_dxa']} DXA: {widths_dxa}"
        )
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    tbl_pr = table._tbl.tblPr

    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(PRESET["content_width_dxa"]))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent_dxa))
    tbl_ind.set(qn("w:type"), "dxa")

    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")

    margins = tbl_pr.find(qn("w:tblCellMar"))
    if margins is None:
        margins = OxmlElement("w:tblCellMar")
        tbl_pr.append(margins)
    for side, value in PRESET["cell_margin_dxa"].items():
        node = margins.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")

    old_grid = table._tbl.tblGrid
    for child in list(old_grid):
        old_grid.remove(child)
    for width in widths_dxa:
        grid_col = OxmlElement("w:gridCol")
        grid_col.set(qn("w:w"), str(width))
        old_grid.append(grid_col)

    for row in table.rows:
        for index, cell in enumerate(row.cells):
            set_cell_width(cell, widths_dxa[index])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_table_borders(table)


def add_table(doc: Document, headers: list[str], rows: list[list[str]], widths_dxa: list[int]):
    table = doc.add_table(rows=1, cols=len(headers))
    for index, text in enumerate(headers):
        cell = table.rows[0].cells[index]
        cell.text = text
        set_cell_shading(cell, PRESET["header_fill"])
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in paragraph.runs:
            set_font(run, size=9, bold=True, color=PRESET["ink"])
    set_repeat_table_header(table.rows[0])
    for row_values in rows:
        cells = table.add_row().cells
        for index, text in enumerate(row_values):
            cells[index].text = str(text)
            paragraph = cells[index].paragraphs[0]
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.space_after = Pt(0)
            for run in paragraph.runs:
                set_font(run, size=9, color="20252B")
    set_table_geometry(table, widths_dxa)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_table_note(doc: Document, text: str):
    paragraph = doc.add_paragraph(style="Table Citation")
    paragraph.add_run(text)
    return paragraph


def add_body(doc: Document, text: str, *, bold_lead: str | None = None):
    paragraph = doc.add_paragraph()
    if bold_lead and text.startswith(bold_lead):
        lead = paragraph.add_run(bold_lead)
        set_font(lead, bold=True, color=PRESET["ink"])
        paragraph.add_run(text[len(bold_lead):])
    else:
        paragraph.add_run(text)
    return paragraph


def add_equation(doc: Document, text: str):
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(4)
    paragraph.paragraph_format.space_after = Pt(8)
    run = paragraph.add_run(text)
    set_font(run, latin="Cambria Math", east_asia="Microsoft YaHei", size=10.5, color=PRESET["ink"])
    return paragraph


def add_figure(doc: Document, filename: str, caption: str, width=6.2):
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.keep_with_next = True
    paragraph.add_run().add_picture(str(FIGURES / filename), width=Inches(width))
    cap = doc.add_paragraph(style="Caption")
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.add_run(caption)
    return cap


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instruction, separate, end])
    set_font(run, size=9, color=PRESET["muted"])


def configure_document(doc: Document):
    section = doc.sections[0]
    section.orientation = WD_ORIENT.PORTRAIT
    section.page_width = Inches(PRESET["page_width"])
    section.page_height = Inches(PRESET["page_height"])
    section.top_margin = Inches(PRESET["margin"])
    section.bottom_margin = Inches(PRESET["margin"])
    section.left_margin = Inches(PRESET["margin"])
    section.right_margin = Inches(PRESET["margin"])
    section.header_distance = Inches(PRESET["header_footer"])
    section.footer_distance = Inches(PRESET["header_footer"])
    section.different_first_page_header_footer = True

    styles = doc.styles
    normal = styles["Normal"]
    set_style_font(normal)
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string("20252B")
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(PRESET["body_after_pt"])
    normal.paragraph_format.line_spacing = PRESET["body_line_spacing"]

    for style_name, token_name in (("Heading 1", "h1"), ("Heading 2", "h2"), ("Heading 3", "h3")):
        style = styles[style_name]
        token = PRESET[token_name]
        set_style_font(style)
        style.font.size = Pt(token["size"])
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(token["color"])
        style.paragraph_format.space_before = Pt(token["before"])
        style.paragraph_format.space_after = Pt(token["after"])
        style.paragraph_format.keep_with_next = True

    title = styles["Title"]
    set_style_font(title)
    title.font.size = Pt(26)
    title.font.bold = True
    title.font.color.rgb = RGBColor.from_string(PRESET["ink"])
    title.paragraph_format.space_before = Pt(0)
    title.paragraph_format.space_after = Pt(12)
    title.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT

    caption = styles["Caption"]
    set_style_font(caption)
    caption.font.size = Pt(9)
    caption.font.color.rgb = RGBColor.from_string(PRESET["muted"])
    caption.paragraph_format.space_before = Pt(4)
    caption.paragraph_format.space_after = Pt(8)
    caption.paragraph_format.keep_with_next = False

    if "Table Citation" not in styles:
        table_note = styles.add_style("Table Citation", WD_STYLE_TYPE.PARAGRAPH)
    else:
        table_note = styles["Table Citation"]
    set_style_font(table_note)
    table_note.font.size = Pt(8.5)
    table_note.font.color.rgb = RGBColor.from_string(PRESET["muted"])
    table_note.paragraph_format.space_before = Pt(4)
    table_note.paragraph_format.space_after = Pt(4)

    header = section.header.paragraphs[0]
    header.text = "2025 CUMCM C题  |  NIPT 时点选择与异常判定"
    header.alignment = WD_ALIGN_PARAGRAPH.LEFT
    for run in header.runs:
        set_font(run, size=8.5, color=PRESET["muted"])
    footer = section.footer.paragraphs[0]
    add_page_number(footer)


def build_report():
    doc = Document()
    configure_document(doc)

    eyebrow = doc.add_paragraph()
    eyebrow.paragraph_format.space_after = Pt(8)
    run = eyebrow.add_run("2025 全国大学生数学建模竞赛 · C题")
    set_font(run, size=10, bold=True, color="2E74B5")

    doc.add_paragraph("NIPT 的时点选择与胎儿异常判定", style="Title")
    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(18)
    run = subtitle.add_run("混合效应 · 区间删失 AFT · 分组交叉验证")
    set_font(run, size=11, color=PRESET["muted"])

    summary = doc.add_table(rows=1, cols=1)
    summary.cell(0, 0).text = "摘要\n本文对 267 名男胎孕妇的 1082 条重复检测记录和 147 名女胎孕妇的 605 条记录进行建模。问题一采用随机截距混合效应模型；问题二、三把首次达到 4% Y 浓度视为区间删失事件，建立对数正态 AFT 模型；问题四按孕妇分组进行五折分类验证。结果显示孕周对 Y 浓度有显著正效应，BMI 有显著负效应；仅 BMI 模型给出 13周3天至19周5天的四档推荐时点；女胎逻辑回归 ROC-AUC 为 0.826。结论仅适用于附件数据和赛题假设，不构成临床建议。"
    set_cell_shading(summary.cell(0, 0), PRESET["callout_fill"])
    set_table_geometry(summary, [9000])
    for index, paragraph in enumerate(summary.cell(0, 0).paragraphs):
        paragraph.paragraph_format.space_after = Pt(4)
        for run in paragraph.runs:
            set_font(run, size=10.5, bold=(index == 0), color=PRESET["ink"] if index == 0 else "20252B")

    keywords = doc.add_paragraph()
    keywords.paragraph_format.space_before = Pt(8)
    keywords.paragraph_format.space_after = Pt(14)
    run = keywords.add_run("关键词  ")
    set_font(run, bold=True, color=PRESET["ink"])
    keywords.add_run("NIPT  ·  混合效应模型  ·  区间删失  ·  BMI 分组  ·  分组交叉验证")

    doc.add_heading("1 题目与数据", level=1)
    add_body(doc, "NIPT 利用母体血液中的胎儿游离 DNA 进行染色体异常筛查。题目要求分析男胎 Y 染色体浓度与孕周、BMI 等因素的关系，为不同 BMI 人群选择检测时点，并对女胎异常进行判定。附件包含同一孕妇的多次采血和重复检测，因此分析与验证均以孕妇代码为分组单位。")
    add_table(
        doc,
        ["工作表", "记录数", "孕妇数", "平均重复", "孕周缺失", "阳性数", "阳性率"],
        [["男胎", "1082", "267", "4.052", "1", "937", "0.866"], ["女胎", "605", "147", "4.116", "0", "67", "0.111"]],
        [1154, 1154, 1154, 1442, 1346, 1250, 1500],
    )
    add_table_note(doc, "表 1  数据概览。女胎阳性标签来自附件“染色体的非整倍体”字段。")

    doc.add_heading("建模约束", level=2)
    for label, text in [
        ("A1", "同一孕妇多次检测共享稳定个体差异，观测之间存在组内相关。"),
        ("A2", "最后一次未达标与首次达标之间包含真实达标时点；首次观测即达标属于左删失。"),
        ("A3", "推荐时点为组内预测达标比例首次达到 90% 的孕周。"),
        ("A4", "时点模型只使用检测前可获得变量，避免检测后信息泄漏。"),
    ]:
        add_body(doc, f"{label}  {text}", bold_lead=label)

    section_heading = doc.add_heading("2 问题一：Y 浓度关系", level=1)
    section_heading.paragraph_format.page_break_before = True
    add_body(doc, "为处理同一孕妇多次检测的组内相关，建立随机截距混合效应模型。响应变量采用 log(Y 浓度)，固定效应包括孕周一次项和二次项、BMI、孕周与 BMI 交互、年龄、GC 含量和过滤读段比例。")
    add_equation(doc, "log Yᵢⱼ = β₀ + β₁Wc + β₂Wc² + β₃Bc + β₄WcBc + β₅Ac + β₆GCc + β₇Fc + uᵢ + εᵢⱼ")
    add_table(
        doc,
        ["指标", "估计", "p 值", "解释"],
        [
            ["孕周一次项", "0.0382", "3.23×10⁻⁴⁸", "显著正效应"],
            ["孕周二次项", "0.00237", "8.35×10⁻⁶", "存在非线性"],
            ["BMI", "-0.0224", "0.00258", "显著负效应"],
            ["孕周×BMI", "0.000664", "0.326", "交互不显著"],
            ["GC 含量", "6.950", "0.00947", "质量因素显著"],
        ],
        [2019, 1442, 1731, 3808],
    )
    add_table_note(doc, "表 2  混合效应模型的关键固定效应。完整系数见 tables/q1_mixed_effects.csv。")
    add_body(doc, "模型成功收敛。ICC=0.7318，边际 R²=0.1440，条件 R²=0.7705，说明孕妇之间的稳定差异是主要变异来源。按孕妇分组的五折 MAE 为 0.0270；组外 R² 略为负值，表明仅凭固定人口变量难以精确预测新孕妇的单次浓度。")
    add_figure(doc, "q1_relationships.png", "图 1  孕周、BMI 与男胎 Y 染色体浓度关系", width=6.15)

    section_heading = doc.add_heading("3 问题二：BMI 分组与时点", level=1)
    section_heading.paragraph_format.page_break_before = True
    add_body(doc, "对每名孕妇构造首次达到 4% 的时间区间 (Lᵢ,Uᵢ]，并建立仅含 BMI 的对数正态 AFT 模型。四个连续 BMI 区间通过动态规划确定，在每组至少 35 人的条件下，最小化组内个体 90% 达标分位数的离差。")
    add_equation(doc, "log Tᵢ = α₀ + α₁BMIᵢ + σεᵢ，εᵢ ~ N(0,1)")
    q2 = read_csv("q2_groups.csv")
    add_table(
        doc,
        ["组", "BMI 区间", "人数", "推荐时点", "达标比例"],
        [[r["组别"], r["BMI区间"], r["孕妇数"], r["推荐时点"], f'{float(r["预测达标比例"]):.4f}'] for r in q2],
        [865, 2500, 1154, 2115, 2366],
    )
    add_table_note(doc, "表 3  问题二的 BMI 分组与推荐检测时点；各组预测达标比例均不低于 0.90。")

    doc.add_heading("4 问题三：多因素修正", level=1)
    add_body(doc, "在 BMI 基础上加入年龄、身高、受孕方式、怀孕次数和生产次数，仍采用区间删失 AFT。检测后才能获得的 Z 值、读段比例和 Y 浓度不进入时点模型。")
    q3 = read_csv("q3_groups.csv")
    add_table(
        doc,
        ["组", "BMI 区间", "人数", "推荐时点", "达标比例"],
        [[r["组别"], r["BMI区间"], r["孕妇数"], r["推荐时点"], f'{float(r["预测达标比例"]):.4f}'] for r in q3],
        [865, 2500, 1154, 2115, 2366],
    )
    add_table_note(doc, "表 4  多因素修正后的分组建议。")
    add_body(doc, "多因素模型 AIC=364.38，高于仅 BMI 模型的 358.57。因此，多因素路线用于刻画异质性和风险修正，不宣称其拟合优于 BMI 主模型。")
    add_figure(doc, "q2_q3_reach_probability.png", "图 2  两条路线下各 BMI 组的预测达标概率", width=6.15)

    doc.add_heading("5 检测误差与稳健性", level=1)
    add_body(doc, "将达标阈值分别设为 3.5%、4.0% 和 4.5%，每次重新构造删失区间并重新拟合。相对 4.0% 基准，推荐时点最多变化 17 天；高 BMI 组对阈值误差更敏感。")
    add_figure(doc, "q2_q3_error_sensitivity.png", "图 3  推荐时点相对 4.0% 基准的变化天数", width=6.15)
    add_body(doc, "删失构成为：左删失 218 人、区间删失 42 人、右删失 7 人。直接把首次观测周当作达标周会系统性推迟估计，区间删失模型更符合附件的观测机制，但结果仍依赖对数正态分布假设。")

    doc.add_heading("6 问题四：女胎异常判定", level=1)
    add_body(doc, "以 Z13、Z18、Z21、ZX、X 浓度、GC 指标、读段数及比例、BMI、年龄和孕周为输入，比较逻辑回归、随机森林、直方图梯度提升和传统 Z 阈值。五折划分始终保证同一孕妇只出现在一个折中。")
    metrics = read_csv("q4_model_metrics.csv")
    add_table(
        doc,
        ["模型", "ROC-AUC", "PR-AUC", "召回率", "特异度", "F1"],
        [[r["模型"], f'{float(r["roc_auc"]):.3f}', f'{float(r["pr_auc"]):.3f}', f'{float(r["recall"]):.3f}', f'{float(r["specificity"]):.3f}', f'{float(r["f1"]):.3f}'] for r in metrics],
        [2560, 1288, 1288, 1288, 1288, 1288],
    )
    add_table_note(doc, "表 5  按孕妇分组五折交叉验证。阳性率较低，PR-AUC 比准确率更有解释价值。")
    add_body(doc, "分组逻辑回归的排序能力最好且可解释性强。阈值 0.5 下，TN=413、FP=125、FN=17、TP=50，对应召回率 0.746、特异度 0.768。传统 |Z|≥3 在本附件上表现较弱，说明异常标签不能由单次三条常染色体 Z 值完全复现。")
    add_figure(doc, "q4_roc_pr.png", "图 4  女胎异常判定的 ROC 与 PR 曲线", width=6.15)
    calibration = read_csv("q4_calibration_summary.csv")[0]
    add_body(
        doc,
        f'类别加权输出只作为筛查分数。嵌套孕妇分组 Platt 校准将 Brier 分数由 {float(calibration["原始Brier"]):.3f} 降至 {float(calibration["校准后Brier"]):.3f}，ECE={float(calibration["ECE"]):.3f}；校准仍需独立样本复核。',
    )
    add_figure(doc, "q4_calibration_threshold.png", "图 5  主模型概率校准与阈值权衡", width=6.15)
    add_figure(doc, "q4_confusion_coefficients.png", "图 6  主模型混淆矩阵与标准化系数", width=5.25)

    doc.add_heading("7 结论与边界", level=1)
    add_body(doc, "问题一：孕周增加与 Y 浓度上升显著相关，BMI 增加与 Y 浓度下降显著相关；ICC=0.732 说明个体差异不可忽略。", bold_lead="问题一：")
    add_body(doc, "问题二：仅 BMI 路线给出四档时点 13周3天、14周4天、16周2天和19周5天。", bold_lead="问题二：")
    add_body(doc, "问题三：多因素路线给出 13周6天、15周2天、16周3天和19周5天；其 AIC 更高，因此只作为异质性修正。", bold_lead="问题三：")
    add_body(doc, "问题四：分组逻辑回归 ROC-AUC=0.826、PR-AUC=0.479，优于传统 Z 阈值。", bold_lead="问题四：")

    doc.add_heading("适用范围", level=2)
    add_body(doc, "附件样本以高 BMI 孕妇为主，不能直接外推到一般孕妇。大量首次观测即达标使左删失占比较高，女胎阳性仅 67 条，均需要外部样本重新校准。模型输出只用于数学建模竞赛，不构成医学诊断或临床建议。")

    doc.add_heading("参考文献", level=2)
    references = [
        "[1] 全国大学生数学建模竞赛组委会. 2025 年高教社杯全国大学生数学建模竞赛赛题, 2025.",
        "[2] Deng C, et al. Maternal and fetal factors influencing fetal fraction. 2023. PMID: 37114008.",
        "[3] Gazdarica J, et al. Insights into non-informative results from non-invasive prenatal screening. 2024. PMID: 38452118.",
        "[4] Rolnik DL, et al. Influence of Body Mass Index on Fetal Fraction Increase With Gestation and Cell-Free DNA Test Failure. 2018. PMID: 29995742.",
        "[5] Zhang X, et al. Evaluation of the Z-score accuracy of noninvasive prenatal testing for fetal trisomies 13, 18 and 21. 2021. PMID: 33480032.",
    ]
    for ref in references:
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.left_indent = Inches(0.22)
        paragraph.paragraph_format.first_line_indent = Inches(-0.22)
        paragraph.paragraph_format.space_after = Pt(2)
        paragraph.paragraph_format.line_spacing = 1.0
        run = paragraph.add_run(ref)
        set_font(run, size=9.2, color="20252B")

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build_report()
