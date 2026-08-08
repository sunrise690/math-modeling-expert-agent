"""Run the reproducible Q3--Q5 solvers and fill the official workbooks.

Example
-------
    python -B src/solve_q3_q5.py --seed 20250808 \
        --q3-seed 20250810 --q4-seed 20250809 --q5-seed 20250808

The official templates under ``source/`` are never edited.  Each run copies
them to ``outputs/`` and fills only the designated data cells.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import platform
import shutil
import sys
from pathlib import Path
from typing import Iterable, Mapping, Sequence

import numpy as np
import openpyxl
import scipy
from openpyxl import load_workbook

try:
    from run_identity import RUN_ID
except ImportError:  # pragma: no cover - package-style execution
    from .run_identity import RUN_ID  # type: ignore

try:
    from geometry_engine import (
        DEFAULT_TARGET,
        GRAVITY,
        SMOKE_LIFETIME,
        SMOKE_RADIUS,
        SMOKE_SINK_SPEED,
        CylindricalOcclusionEvaluator,
        interval_duration,
    )
    from optimizer import (
        DEFAULT_RANDOM_SEED,
        MISSILES,
        UAV_INITIALS,
        BombPlan,
        SolveResult,
        exact_evaluate,
        result_to_dict,
        solve_q3,
        solve_q4,
        solve_q5,
        strategy_constraint_violations,
    )
except ImportError:
    from .geometry_engine import (  # type: ignore
        DEFAULT_TARGET,
        GRAVITY,
        SMOKE_LIFETIME,
        SMOKE_RADIUS,
        SMOKE_SINK_SPEED,
        CylindricalOcclusionEvaluator,
        interval_duration,
    )
    from .optimizer import (  # type: ignore
        DEFAULT_RANDOM_SEED,
        MISSILES,
        UAV_INITIALS,
        BombPlan,
        SolveResult,
        exact_evaluate,
        result_to_dict,
        solve_q3,
        solve_q4,
        solve_q5,
        strategy_constraint_violations,
    )


PROJECT_ROOT = Path(__file__).resolve().parents[1]
SOURCE_DIR = PROJECT_ROOT / "source"
OUTPUT_DIR = PROJECT_ROOT / "outputs"
VALIDATION_PATH = PROJECT_ROOT / "validation" / "q3_q5_independent.json"


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(1 << 20), b""):
            digest.update(block)
    return digest.hexdigest()


def _duration_for_plan(result: SolveResult, plan: BombPlan) -> float:
    for candidate, intervals in zip(result.plans, result.exact_individual_intervals):
        if candidate == plan:
            return interval_duration(intervals)
    raise KeyError("plan is not present in the solve result")


def _copy_template(name: str) -> Path:
    source = SOURCE_DIR / name
    destination = OUTPUT_DIR / name
    if not source.exists():
        raise FileNotFoundError(source)
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    shutil.copy2(source, destination)
    return destination


def _fill_result1(result: SolveResult) -> Path:
    path = _copy_template("result1.xlsx")
    workbook = load_workbook(path)
    sheet = workbook[workbook.sheetnames[0]]
    plans = sorted(result.plans, key=lambda plan: plan.release_time)
    if len(plans) != 3 or any(plan.drone_id != "FY1" for plan in plans):
        raise ValueError("Q3 result must contain exactly three FY1 plans")
    for row, plan in zip(range(2, 5), plans):
        release = plan.release_position
        explosion = plan.explosion_position
        values = [
            plan.normalized_heading_deg,
            plan.speed,
            sheet.cell(row, 3).value,
            *release,
            *explosion,
            _duration_for_plan(result, plan),
        ]
        for column, value in enumerate(values, start=1):
            sheet.cell(row, column).value = value
    workbook.save(path)
    return path


def _fill_result2(result: SolveResult) -> Path:
    path = _copy_template("result2.xlsx")
    workbook = load_workbook(path)
    sheet = workbook[workbook.sheetnames[0]]
    by_drone = {plan.drone_id: plan for plan in result.plans}
    for row in range(2, 5):
        drone_id = str(sheet.cell(row, 1).value)
        plan = by_drone[drone_id]
        release = plan.release_position
        explosion = plan.explosion_position
        values = [
            drone_id,
            plan.normalized_heading_deg,
            plan.speed,
            *release,
            *explosion,
            _duration_for_plan(result, plan),
        ]
        for column, value in enumerate(values, start=1):
            sheet.cell(row, column).value = value
    workbook.save(path)
    return path


def _fill_result3(result: SolveResult) -> Path:
    path = _copy_template("result3.xlsx")
    workbook = load_workbook(path)
    sheet = workbook[workbook.sheetnames[0]]
    by_drone: dict[str, list[BombPlan]] = {}
    for plan in result.plans:
        by_drone.setdefault(plan.drone_id, []).append(plan)
    for plans in by_drone.values():
        plans.sort(key=lambda plan: plan.release_time)

    for row in range(2, 17):
        drone_id = str(sheet.cell(row, 1).value)
        bomb_number = int(sheet.cell(row, 4).value)
        candidates = by_drone.get(drone_id, [])
        if bomb_number > len(candidates):
            for column in (2, 3, 5, 6, 7, 8, 9, 10, 11, 12):
                sheet.cell(row, column).value = None
            continue
        plan = candidates[bomb_number - 1]
        release = plan.release_position
        explosion = plan.explosion_position
        values = {
            2: plan.normalized_heading_deg,
            3: plan.speed,
            5: release[0],
            6: release[1],
            7: release[2],
            8: explosion[0],
            9: explosion[1],
            10: explosion[2],
            11: _duration_for_plan(result, plan),
            12: plan.missile_id,
        }
        for column, value in values.items():
            sheet.cell(row, column).value = value
    workbook.save(path)
    return path


def write_workbooks(results: Mapping[str, SolveResult]) -> dict[str, Path]:
    return {
        "Q3": _fill_result1(results["Q3"]),
        "Q4": _fill_result2(results["Q4"]),
        "Q5": _fill_result3(results["Q5"]),
    }


def _full_cylinder_audit(result: SolveResult) -> dict[str, object]:
    """Continuous-time roots with continuously refined cylindrical boundaries."""

    evaluator = CylindricalOcclusionEvaluator(
        DEFAULT_TARGET, n_azimuth=96, n_linear=9, local_starts=4
    )
    individual, union = exact_evaluate(
        result.plans,
        max_step=0.02,
        mode="full_cylinder",
        evaluator=evaluator,
        optimize_full=True,
    )
    return {
        "convention": (
            "continuous-time Brent roots with 96-azimuth/9-linear-node candidate "
            "generation and continuous boundary refinement on the complete cylinder; "
            "conservative secondary audit"
        ),
        "individual_intervals": [
            [list(interval) for interval in intervals] for intervals in individual
        ],
        "individual_durations": [interval_duration(intervals) for intervals in individual],
        "union_by_missile": {
            missile_id: [list(interval) for interval in intervals]
            for missile_id, intervals in union.items()
        },
        "duration_by_missile": {
            missile_id: interval_duration(intervals)
            for missile_id, intervals in union.items()
        },
        "objective": sum(interval_duration(intervals) for intervals in union.values()),
    }


def _parse_result1(
    path: Path,
) -> tuple[
    list[BombPlan],
    list[float],
    list[tuple[tuple[float, float, float], tuple[float, float, float]]],
]:
    workbook = load_workbook(path, data_only=False, read_only=True)
    sheet = workbook[workbook.sheetnames[0]]
    plans: list[BombPlan] = []
    durations: list[float] = []
    table_positions: list[
        tuple[tuple[float, float, float], tuple[float, float, float]]
    ] = []
    initial = np.asarray(UAV_INITIALS["FY1"], dtype=float)
    for row in range(2, 5):
        heading = float(sheet.cell(row, 1).value)
        speed = float(sheet.cell(row, 2).value)
        release_xy = np.array(
            (float(sheet.cell(row, 4).value), float(sheet.cell(row, 5).value))
        )
        explosion_xy = np.array(
            (float(sheet.cell(row, 7).value), float(sheet.cell(row, 8).value))
        )
        table_positions.append(
            (
                (
                    float(sheet.cell(row, 4).value),
                    float(sheet.cell(row, 5).value),
                    float(sheet.cell(row, 6).value),
                ),
                (
                    float(sheet.cell(row, 7).value),
                    float(sheet.cell(row, 8).value),
                    float(sheet.cell(row, 9).value),
                ),
            )
        )
        direction = np.array((np.cos(np.deg2rad(heading)), np.sin(np.deg2rad(heading))))
        release_time = float(np.dot(release_xy - initial[:2], direction) / speed)
        explosion_time = float(np.dot(explosion_xy - initial[:2], direction) / speed)
        plans.append(
            BombPlan("FY1", heading, speed, release_time, explosion_time - release_time, "M1")
        )
        durations.append(float(sheet.cell(row, 10).value))
    return plans, durations, table_positions


def _parse_result2(
    path: Path,
) -> tuple[
    list[BombPlan],
    list[float],
    list[tuple[tuple[float, float, float], tuple[float, float, float]]],
]:
    workbook = load_workbook(path, data_only=False, read_only=True)
    sheet = workbook[workbook.sheetnames[0]]
    plans: list[BombPlan] = []
    durations: list[float] = []
    table_positions: list[
        tuple[tuple[float, float, float], tuple[float, float, float]]
    ] = []
    for row in range(2, 5):
        drone_id = str(sheet.cell(row, 1).value)
        heading = float(sheet.cell(row, 2).value)
        speed = float(sheet.cell(row, 3).value)
        initial = np.asarray(UAV_INITIALS[drone_id], dtype=float)
        release_xy = np.array(
            (float(sheet.cell(row, 4).value), float(sheet.cell(row, 5).value))
        )
        explosion_xy = np.array(
            (float(sheet.cell(row, 7).value), float(sheet.cell(row, 8).value))
        )
        table_positions.append(
            (
                (
                    float(sheet.cell(row, 4).value),
                    float(sheet.cell(row, 5).value),
                    float(sheet.cell(row, 6).value),
                ),
                (
                    float(sheet.cell(row, 7).value),
                    float(sheet.cell(row, 8).value),
                    float(sheet.cell(row, 9).value),
                ),
            )
        )
        direction = np.array((np.cos(np.deg2rad(heading)), np.sin(np.deg2rad(heading))))
        release_time = float(np.dot(release_xy - initial[:2], direction) / speed)
        explosion_time = float(np.dot(explosion_xy - initial[:2], direction) / speed)
        plans.append(
            BombPlan(drone_id, heading, speed, release_time, explosion_time - release_time, "M1")
        )
        durations.append(float(sheet.cell(row, 10).value))
    return plans, durations, table_positions


def _parse_result3(
    path: Path,
) -> tuple[
    list[BombPlan],
    list[float],
    list[tuple[tuple[float, float, float], tuple[float, float, float]]],
]:
    workbook = load_workbook(path, data_only=False, read_only=True)
    sheet = workbook[workbook.sheetnames[0]]
    plans: list[BombPlan] = []
    durations: list[float] = []
    table_positions: list[
        tuple[tuple[float, float, float], tuple[float, float, float]]
    ] = []
    for row in range(2, 17):
        if sheet.cell(row, 2).value is None:
            continue
        drone_id = str(sheet.cell(row, 1).value)
        heading = float(sheet.cell(row, 2).value)
        speed = float(sheet.cell(row, 3).value)
        initial = np.asarray(UAV_INITIALS[drone_id], dtype=float)
        release_xy = np.array(
            (float(sheet.cell(row, 5).value), float(sheet.cell(row, 6).value))
        )
        explosion_xy = np.array(
            (float(sheet.cell(row, 8).value), float(sheet.cell(row, 9).value))
        )
        table_positions.append(
            (
                (
                    float(sheet.cell(row, 5).value),
                    float(sheet.cell(row, 6).value),
                    float(sheet.cell(row, 7).value),
                ),
                (
                    float(sheet.cell(row, 8).value),
                    float(sheet.cell(row, 9).value),
                    float(sheet.cell(row, 10).value),
                ),
            )
        )
        direction = np.array((np.cos(np.deg2rad(heading)), np.sin(np.deg2rad(heading))))
        release_time = float(np.dot(release_xy - initial[:2], direction) / speed)
        explosion_time = float(np.dot(explosion_xy - initial[:2], direction) / speed)
        plans.append(
            BombPlan(
                drone_id,
                heading,
                speed,
                release_time,
                explosion_time - release_time,
                str(sheet.cell(row, 12).value),
            )
        )
        durations.append(float(sheet.cell(row, 11).value))
    return plans, durations, table_positions


def roundtrip_validate_workbook(
    problem: str, path: Path, reference: SolveResult
) -> dict[str, object]:
    parser = {"Q3": _parse_result1, "Q4": _parse_result2, "Q5": _parse_result3}[problem]
    plans, table_durations, table_positions = parser(path)
    individual, union = exact_evaluate(plans, max_step=0.01)
    recomputed = [interval_duration(intervals) for intervals in individual]
    violations = strategy_constraint_violations(plans)
    coordinate_errors: list[float] = []
    reference_plans = sorted(
        reference.plans, key=lambda plan: (plan.drone_id, plan.release_time)
    )
    parsed_plans = sorted(plans, key=lambda plan: (plan.drone_id, plan.release_time))
    for parsed, expected in zip(parsed_plans, reference_plans):
        coordinate_errors.extend(
            abs(a - b)
            for a, b in zip(parsed.release_position, expected.release_position)
        )
        coordinate_errors.extend(
            abs(a - b)
            for a, b in zip(parsed.explosion_position, expected.explosion_position)
        )
    # This second comparison includes the table's z columns; the first
    # comparison primarily checks recovery of absolute times from xy fields.
    for parsed, (table_release, table_explosion) in zip(plans, table_positions):
        coordinate_errors.extend(
            abs(a - b) for a, b in zip(parsed.release_position, table_release)
        )
        coordinate_errors.extend(
            abs(a - b) for a, b in zip(parsed.explosion_position, table_explosion)
        )
    return {
        "rows_parsed": len(plans),
        "constraint_violations": violations,
        "max_constraint_violation": max(violations.values(), default=0.0),
        "max_reconstructed_coordinate_error_m": max(coordinate_errors, default=0.0),
        "max_row_duration_error_s": max(
            (abs(a - b) for a, b in zip(table_durations, recomputed)), default=0.0
        ),
        "recomputed_objective": sum(interval_duration(value) for value in union.values()),
        "reference_objective": reference.exact_objective,
        "objective_error_s": abs(
            sum(interval_duration(value) for value in union.values())
            - reference.exact_objective
        ),
    }


def resolve_problem_seeds(
    seed: int,
    *,
    q3_seed: int | None = None,
    q4_seed: int | None = None,
    q5_seed: int | None = None,
) -> dict[str, int]:
    """Resolve per-problem seeds while retaining ``--seed`` compatibility."""

    return {
        "Q3": seed if q3_seed is None else q3_seed,
        "Q4": seed if q4_seed is None else q4_seed,
        "Q5": seed if q5_seed is None else q5_seed,
    }


def run_all(
    *,
    seed: int,
    quick: bool,
    full_cylinder_audit: bool,
    q3_seed: int | None = None,
    q4_seed: int | None = None,
    q5_seed: int | None = None,
) -> tuple[dict[str, SolveResult], dict[str, object]]:
    problem_seeds = resolve_problem_seeds(
        seed,
        q3_seed=q3_seed,
        q4_seed=q4_seed,
        q5_seed=q5_seed,
    )
    print(f"[Q3] solving (seed={problem_seeds['Q3']}, quick={quick})", flush=True)
    q3 = solve_q3(seed=problem_seeds["Q3"], quick=quick)
    print(f"[Q3] exact objective={q3.exact_objective:.9f}", flush=True)
    print(f"[Q4] solving (seed={problem_seeds['Q4']}, quick={quick})", flush=True)
    q4 = solve_q4(seed=problem_seeds["Q4"], quick=quick)
    print(f"[Q4] exact objective={q4.exact_objective:.9f}", flush=True)
    print(f"[Q5] solving (seed={problem_seeds['Q5']}, quick={quick})", flush=True)
    q5 = solve_q5(seed=problem_seeds["Q5"], quick=quick)
    print(f"[Q5] exact objective={q5.exact_objective:.9f}", flush=True)
    results = {"Q3": q3, "Q4": q4, "Q5": q5}

    workbooks = write_workbooks(results)
    roundtrip = {
        problem: roundtrip_validate_workbook(problem, workbooks[problem], result)
        for problem, result in results.items()
    }
    cylinder = (
        {
            problem: _full_cylinder_audit(result)
            for problem, result in results.items()
        }
        if full_cylinder_audit
        else {}
    )
    source_files = [
        SOURCE_DIR / "A题.pdf",
        SOURCE_DIR / "result1.xlsx",
        SOURCE_DIR / "result2.xlsx",
        SOURCE_DIR / "result3.xlsx",
    ]
    report: dict[str, object] = {
        "schema_version": "1.1",
        "run_id": RUN_ID,
        "scope": "Independent reproducible Q3-Q5 solver using only official source files and geometry_engine.py.",
        "random_seed": seed,
        "random_seed_by_problem": problem_seeds,
        "quick_mode": quick,
        "model": {
            "target_centerline_point": list(DEFAULT_TARGET.center),
            "target_bottom_center": list(DEFAULT_TARGET.bottom_center),
            "target_radius_m": DEFAULT_TARGET.radius,
            "target_height_m": DEFAULT_TARGET.height,
            "gravity_m_s2": GRAVITY,
            "smoke_radius_m": SMOKE_RADIUS,
            "smoke_lifetime_s": SMOKE_LIFETIME,
            "smoke_sink_speed_m_s": SMOKE_SINK_SPEED,
            "missile_impact_times_s": {
                missile_id: missile.impact_time for missile_id, missile in MISSILES.items()
            },
        },
        "source_sha256": {path.name: sha256_file(path) for path in source_files},
        "runtime": {
            "python": platform.python_version(),
            "numpy": np.__version__,
            "scipy": scipy.__version__,
            "openpyxl": openpyxl.__version__,
        },
        "results": {problem: result_to_dict(result) for problem, result in results.items()},
        "full_cylinder_secondary_audit": cylinder,
        "workbook_roundtrip_validation": roundtrip,
        "output_workbook_sha256": {
            workbooks[problem].name: sha256_file(workbooks[problem])
            for problem in workbooks
        },
        "reproduction_command": f"python -B src/solve_q3_q5.py --seed {seed}"
        f" --q3-seed {problem_seeds['Q3']} --q4-seed {problem_seeds['Q4']}"
        f" --q5-seed {problem_seeds['Q5']}"
        + (" --quick" if quick else ""),
        "optimality_disclaimer": (
            "All reported strategies are continuously verified feasible solutions. "
            "Q5 is a restricted route-library/beam-search solution and is not claimed globally optimal."
        ),
    }
    return results, report


def parse_args(argv: Sequence[str] | None = None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--seed",
        type=int,
        default=DEFAULT_RANDOM_SEED,
        help="Compatibility default used by any problem without an explicit per-problem seed.",
    )
    parser.add_argument("--q3-seed", type=int, default=None)
    parser.add_argument("--q4-seed", type=int, default=None)
    parser.add_argument("--q5-seed", type=int, default=None)
    parser.add_argument(
        "--quick", action="store_true", help="Use smaller route grids for a smoke test."
    )
    parser.add_argument(
        "--skip-full-cylinder-audit",
        action="store_true",
        help="Skip the slower continuously refined complete-cylinder secondary audit.",
    )
    return parser.parse_args(argv)


def main(argv: Sequence[str] | None = None) -> int:
    args = parse_args(argv)
    _results, report = run_all(
        seed=args.seed,
        q3_seed=args.q3_seed,
        q4_seed=args.q4_seed,
        q5_seed=args.q5_seed,
        quick=args.quick,
        full_cylinder_audit=not args.skip_full_cylinder_audit,
    )
    VALIDATION_PATH.parent.mkdir(parents=True, exist_ok=True)
    VALIDATION_PATH.write_text(
        json.dumps(report, ensure_ascii=False, indent=2, sort_keys=True) + "\n",
        encoding="utf-8",
    )
    print(f"validation: {VALIDATION_PATH}", flush=True)
    return 0


if __name__ == "__main__":
    raise SystemExit(main(sys.argv[1:]))
