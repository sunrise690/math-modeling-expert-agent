"""Build the paper appendix from checked source code and result workbooks.

The final appendix is text-first: LaTeX fragments contain searchable source
excerpts and three-line-table rows read directly from the submitted XLSX
files.  Four plain SVG previews are also emitted for the project's existing
evidence-manifest gate; they are QA aids and are not embedded in the paper.
No screenshot styling, decorative cards, or synthetic data are used.
"""

from __future__ import annotations

import hashlib
import html
import importlib.metadata
import json
import sys
from pathlib import Path
from typing import Any, Iterable, Sequence

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter


ROOT = Path(__file__).resolve().parents[1]
OUTPUT_DIR = ROOT / "support" / "appendix-evidence"
MANIFEST_PATH = OUTPUT_DIR / "manifest.json"

# These four bitmaps belonged to the former poster-style appendix.  They are
# generated artifacts, not source material; the text-first renderer removes
# them once their native-text replacements have been written.
LEGACY_POSTERS = (
    "appendix_code_geometry.png",
    "appendix_code_optimizer.png",
    "appendix_data_result12.png",
    "appendix_data_result3.png",
)

GEOMETRY = ROOT / "src" / "geometry_engine.py"
OPTIMIZER = ROOT / "src" / "optimizer.py"
SOLVE_Q3_Q5 = ROOT / "src" / "solve_q3_q5.py"

FILE_INDEX: tuple[tuple[str, Path, str], ...] = (
    ("连续几何与区间运算", GEOMETRY, "527--704"),
    ("问题一、二全局搜索与精化", ROOT / "src" / "solve_q1_q2.py", "530--896"),
    ("问题三至五组合优化", OPTIMIZER, "470--640，790--1203"),
    ("结果表写入与反算", SOLVE_Q3_Q5, "112--204，397--442"),
    ("问题二多种子试验", ROOT / "src" / "validate_q2_multiseed.py", "全文"),
    ("问题三、四多种子试验", ROOT / "src" / "validate_q3_q4_multiseed.py", "全文"),
    ("一键核验入口", ROOT / "src" / "run_all.py", "1142--1856，1914--2280，2826--2888"),
)

CODE_FRAGMENT_SPECS: dict[str, list[tuple[str, Path, int, int]]] = {
    "appendix_code_geometry.tex": [
        ("程序 B1.1\quad 连续边界定位", GEOMETRY, 553, 581),
        ("程序 B1.2\quad 时间区间并集", GEOMETRY, 684, 704),
    ],
    "appendix_code_evaluation.tex": [
        ("程序 B2.1\quad 同机硬约束复核", OPTIMIZER, 589, 605),
        ("程序 B2.2\quad 连续区间精确复算", OPTIMIZER, 619, 640),
    ],
    "appendix_code_q5.tex": [
        ("程序 B3.1\quad 多无人机有限状态保留", OPTIMIZER, 1140, 1164),
        ("程序 B3.2\quad 候选方案连续复算与重排", OPTIMIZER, 1185, 1203),
    ],
}


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def _sha256_text(text: str) -> str:
    return hashlib.sha256(text.encode("utf-8")).hexdigest()


def _relative(path: Path) -> str:
    return path.relative_to(ROOT).as_posix()


def _tex_escape(value: object) -> str:
    text = str(value)
    replacements = {
        "\\": r"\textbackslash{}",
        "&": r"\&",
        "%": r"\%",
        "$": r"\$",
        "#": r"\#",
        "_": r"\_",
        "{": r"\{",
        "}": r"\}",
        "~": r"\textasciitilde{}",
        "^": r"\textasciicircum{}",
    }
    return "".join(replacements.get(character, character) for character in text)


def _format_number(value: object, digits: int = 4) -> str:
    number = float(value)
    if abs(number) < 0.5 * 10 ** (-digits):
        number = 0.0
    return f"{number:.{digits}f}"


def _format_scientific(value: float, digits: int = 3) -> str:
    if value == 0.0:
        return "$0$"
    mantissa, exponent = f"{value:.{digits}e}".split("e")
    return f"${mantissa}\\times 10^{{{int(exponent)}}}$"


def _tex_row(values: Iterable[str]) -> str:
    return " & ".join(values) + r" \\"


def _read_active_rows(path: Path) -> tuple[str, list[list[object]]]:
    workbook = load_workbook(path, data_only=True, read_only=True)
    worksheet = workbook[workbook.sheetnames[0]]
    raw_rows = [list(row) for row in worksheet.iter_rows(values_only=True)]
    active: list[list[object]] = []
    for row in raw_rows:
        if not any(value is not None and str(value).strip() for value in row):
            if active:
                break
            continue
        active.append(row)
    if not active:
        raise RuntimeError(f"Workbook is empty: {path}")
    return worksheet.title, active


def _workbook_snapshot(path: Path) -> dict[str, Any]:
    sheet, rows = _read_active_rows(path)
    column_count = len(rows[0])
    normalized_rows = [
        ["" if value is None else str(value) for value in row] for row in rows
    ]
    normalized = json.dumps(
        normalized_rows, ensure_ascii=False, separators=(",", ":")
    )
    return {
        "path": _relative(path),
        "sha256": _sha256(path),
        "selection": f"{sheet}!A1:{get_column_letter(column_count)}{len(rows)}",
        "active_row_count_including_header": len(rows),
        "column_count": column_count,
        "normalized_value_sha256": _sha256_text(normalized),
    }


def _code_snapshot(path: Path, start: int, end: int) -> dict[str, Any]:
    lines = path.read_text(encoding="utf-8").splitlines()
    excerpt = "\n".join(lines[start - 1 : end]) + "\n"
    return {
        "path": _relative(path),
        "sha256": _sha256(path),
        "selection": {"kind": "line_range", "start": start, "end": end},
        "normalized_excerpt_sha256": _sha256_text(excerpt),
    }


def _write_code_fragments() -> list[dict[str, Any]]:
    records: list[dict[str, Any]] = []
    for filename, snippets in CODE_FRAGMENT_SPECS.items():
        output = OUTPUT_DIR / filename
        parts = ["% Generated by src/render_appendix_evidence.py; do not edit.\n"]
        sources: list[dict[str, Any]] = []
        maximum_width = 0
        for index, (title, source, start, end) in enumerate(snippets):
            lines = source.read_text(encoding="utf-8").splitlines()[start - 1 : end]
            numbered = [f"{line_number:4d}  {line}" for line_number, line in enumerate(lines, start)]
            maximum_width = max(maximum_width, *(len(line) for line in numbered))
            if index:
                parts.append("\\medskip\n")
            parts.extend(
                [
                    f"\\noindent\\textbf{{{title}}}\\hfill"
                    f"{{\\footnotesize\\texttt{{{_tex_escape(_relative(source))}:{start}--{end}}}}}\\par\n",
                    f"\\noindent{{\\scriptsize SHA-256: \\texttt{{{_sha256(source)[:16]}}}}}\\par\n",
                    "\\begingroup\n",
                    "\\fontsize{6.5pt}{7.45pt}\\selectfont\n",
                    "\\begin{verbatim}\n",
                    "\n".join(numbered) + "\n",
                    "\\end{verbatim}\n",
                    "\\endgroup\n",
                ]
            )
            sources.append(_code_snapshot(source, start, end))
        if maximum_width > 122:
            raise RuntimeError(f"Code line is too wide for appendix: {filename} ({maximum_width})")
        output.write_text("".join(parts), encoding="utf-8")
        records.append(
            {
                "path": _relative(output),
                "sha256": _sha256(output),
                "kind": "searchable_latex_source_excerpt",
                "maximum_display_characters": maximum_width,
                "sources": sources,
            }
        )
    return records


def _write_file_index() -> dict[str, Any]:
    output = OUTPUT_DIR / "appendix_file_rows.tex"
    rows = [
        _tex_row(
            [
                purpose,
                f"\\texttt{{{_tex_escape(_relative(path))}}}",
                location,
                f"\\texttt{{{_sha256(path)[:12]}}}",
            ]
        )
        for purpose, path, location in FILE_INDEX
    ]
    output.write_text(
        "% Generated from the current source tree; do not edit.\n"
        + "\n".join(rows)
        + "\n\\bottomrule\n",
        encoding="utf-8",
    )
    return {
        "path": _relative(output),
        "sha256": _sha256(output),
        "kind": "searchable_latex_file_index",
        "sources": [
            {"path": _relative(path), "sha256": _sha256(path)}
            for _purpose, path, _location in FILE_INDEX
        ],
    }


def _coordinate(row: Sequence[object], indices: Sequence[int]) -> str:
    values = [f"{_format_number(row[index])}" for index in indices]
    return "$(" + ",\\,".join(values) + ")$"


def _write_table_rows() -> tuple[list[dict[str, Any]], dict[str, Any]]:
    result1 = ROOT / "outputs" / "result1.xlsx"
    result2 = ROOT / "outputs" / "result2.xlsx"
    result3 = ROOT / "outputs" / "result3.xlsx"
    workbooks = (result1, result2, result3)
    active = {path.name: _read_active_rows(path)[1] for path in workbooks}
    expected_counts = {"result1.xlsx": 4, "result2.xlsx": 4, "result3.xlsx": 16}
    actual_counts = {name: len(rows) for name, rows in active.items()}
    if actual_counts != expected_counts:
        raise RuntimeError(f"Unexpected workbook active ranges: {actual_counts}")

    table_specs: list[tuple[str, Path, list[str]]] = []

    result1_rows = [
        _tex_row(
            [
                _tex_escape(int(row[2])),
                _format_number(row[0]),
                _format_number(row[1]),
                _coordinate(row, (3, 4, 5)),
                _coordinate(row, (6, 7, 8)),
                _format_number(row[9], 6),
            ]
        )
        for row in active[result1.name][1:]
    ]
    table_specs.append(("appendix_result1_rows.tex", result1, result1_rows))

    result2_rows = [
        _tex_row(
            [
                _tex_escape(row[0]),
                _format_number(row[1]),
                _format_number(row[2]),
                _coordinate(row, (3, 4, 5)),
                _coordinate(row, (6, 7, 8)),
                _format_number(row[9], 6),
            ]
        )
        for row in active[result2.name][1:]
    ]
    table_specs.append(("appendix_result2_rows.tex", result2, result2_rows))

    result3_rows = [
        _tex_row(
            [
                _tex_escape(row[0]),
                _tex_escape(int(row[3])),
                _tex_escape(row[11]),
                _format_number(row[1]),
                _format_number(row[2]),
                _coordinate(row, (4, 5, 6)),
                _coordinate(row, (7, 8, 9)),
                _format_number(row[10], 6),
            ]
        )
        for row in active[result3.name][1:]
    ]
    table_specs.append(("appendix_result3_rows.tex", result3, result3_rows))

    records: list[dict[str, Any]] = []
    for filename, source, rows in table_specs:
        output = OUTPUT_DIR / filename
        output.write_text(
            "% Generated from " + _relative(source) + "; do not edit.\n"
            + "\n".join(rows)
            + "\n\\bottomrule\n",
            encoding="utf-8",
        )
        records.append(
            {
                "path": _relative(output),
                "sha256": _sha256(output),
                "kind": "searchable_latex_table_rows",
                "sources": [_workbook_snapshot(source)],
            }
        )

    durations: dict[str, float] = {}
    row_counts: dict[str, int] = {}
    duration_columns = {"result1.xlsx": 9, "result2.xlsx": 9, "result3.xlsx": 10}
    for name, rows in active.items():
        row_counts[name] = len(rows) - 1
        durations[name] = sum(float(row[duration_columns[name]]) for row in rows[1:])

    q35_path = ROOT / "validation" / "q3_q5_independent.json"
    q35 = json.loads(q35_path.read_text(encoding="utf-8"))
    union_rows = []
    for workbook_name, question, note in (
        ("result1.xlsx", "Q3", "存在区间重叠，按并集计时"),
        ("result2.xlsx", "Q4", "三段互不重叠"),
        ("result3.xlsx", "Q5", "按导弹分别取并集后求和"),
    ):
        union = float(q35["results"][question]["exact_centerline_objective"])
        union_rows.append(
            _tex_row(
                [
                    f"\\texttt{{{workbook_name}}}",
                    question,
                    str(row_counts[workbook_name]),
                    _format_number(durations[workbook_name], 6),
                    _format_number(union, 6),
                    note,
                ]
            )
        )
    union_output = OUTPUT_DIR / "appendix_union_rows.tex"
    union_output.write_text(
        "% Generated from result workbooks and validation JSON; do not edit.\n"
        + "\n".join(union_rows)
        + "\n\\bottomrule\n",
        encoding="utf-8",
    )
    records.append(
        {
            "path": _relative(union_output),
            "sha256": _sha256(union_output),
            "kind": "searchable_latex_aggregate_rows",
            "sources": [
                *[_workbook_snapshot(path) for path in workbooks],
                {"path": _relative(q35_path), "sha256": _sha256(q35_path)},
            ],
        }
    )
    return records, {"row_counts": row_counts, "duration_sums": durations}


def _load_json(path: Path) -> dict[str, Any]:
    return json.loads(path.read_text(encoding="utf-8"))


def _write_validation_rows() -> tuple[list[dict[str, Any]], dict[str, Any]]:
    q12_path = ROOT / "validation" / "q1_q2_independent.json"
    q2_multi_path = ROOT / "validation" / "q2_multiseed.json"
    q35_path = ROOT / "validation" / "q3_q5_independent.json"
    q34_multi_path = ROOT / "validation" / "q3_q4_multiseed.json"
    q12 = _load_json(q12_path)
    q2_multi = _load_json(q2_multi_path)
    q35 = _load_json(q35_path)
    q34_multi = _load_json(q34_multi_path)

    center = {
        "Q1": float(q12["q1"]["g_9.80"]["centerline"]["duration_s"]),
        "Q2": float(q12["q2"]["centerline"]["duration_s"]),
        **{
            question: float(q35["results"][question]["exact_centerline_objective"])
            for question in ("Q3", "Q4", "Q5")
        },
    }
    conservative = {
        "Q1": float(q12["q1"]["g_9.80"]["full_cylinder"]["duration_s"]),
        "Q2": float(q12["q2"]["full_cylinder_tau_0_boundary"]["duration_s"]),
        **{
            question: float(q35["full_cylinder_secondary_audit"][question]["objective"])
            for question in ("Q3", "Q4", "Q5")
        },
    }
    q2_summary = q2_multi["summary"]
    q3_summary = q34_multi["problems"]["Q3"]["summary"]
    q4_summary = q34_multi["problems"]["Q4"]["summary"]
    repeat_evidence = {
        "Q1": "题设给定方案",
        "Q2": f"8 组，极差 {_format_scientific(float(q2_summary['maximum_s']) - float(q2_summary['minimum_s']))}",
        "Q3": f"5 组，极差 {_format_number(float(q3_summary['maximum_s']) - float(q3_summary['minimum_s']), 6)} s",
        "Q4": f"5 组，极差 {_format_number(float(q4_summary['maximum_s']) - float(q4_summary['minimum_s']), 6)} s",
        "Q5": "固定种子；无多种子证据",
    }
    interpretations = {
        "Q1": "给定方案复算",
        "Q2": "重复搜索稳定",
        "Q3": "可行解，受初值影响",
        "Q4": "可行解，重复性较好",
        "Q5": "受限路线库可行解",
    }
    objective_rows = [
        _tex_row(
            [
                question,
                _format_number(center[question], 6),
                _format_number(conservative[question], 6),
                repeat_evidence[question],
                interpretations[question],
            ]
        )
        for question in ("Q1", "Q2", "Q3", "Q4", "Q5")
    ]

    roundtrip = q35["workbook_roundtrip_validation"]
    max_constraint = max(
        float(roundtrip[question]["max_constraint_violation"])
        for question in ("Q3", "Q4", "Q5")
    )
    max_coordinate = max(
        float(roundtrip[question]["max_reconstructed_coordinate_error_m"])
        for question in ("Q3", "Q4", "Q5")
    )
    max_duration = max(
        float(roundtrip[question]["max_row_duration_error_s"])
        for question in ("Q3", "Q4", "Q5")
    )
    max_objective = max(
        float(roundtrip[question]["objective_error_s"])
        for question in ("Q3", "Q4", "Q5")
    )
    recheck_rows = [
        _tex_row(
            [
                _format_scientific(max_constraint),
                _format_scientific(max_coordinate),
                _format_scientific(max_duration),
                _format_scientific(max_objective),
                "$10^{-8}$",
                "通过",
            ]
        )
    ]

    validation_specs = [
        ("appendix_objective_rows.tex", objective_rows, [q12_path, q2_multi_path, q35_path, q34_multi_path]),
        ("appendix_recheck_rows.tex", recheck_rows, [q35_path]),
    ]
    records: list[dict[str, Any]] = []
    for filename, rows, sources in validation_specs:
        output = OUTPUT_DIR / filename
        output.write_text(
            "% Generated from validation JSON; do not edit.\n"
            + "\n".join(rows)
            + "\n\\bottomrule\n",
            encoding="utf-8",
        )
        records.append(
            {
                "path": _relative(output),
                "sha256": _sha256(output),
                "kind": "searchable_latex_validation_rows",
                "sources": [
                    {"path": _relative(source), "sha256": _sha256(source)}
                    for source in sources
                ],
            }
        )
    return records, {
        "centerline_objectives_s": center,
        "full_cylinder_objectives_s": conservative,
        "roundtrip_maxima": {
            "constraint_violation": max_constraint,
            "coordinate_error_m": max_coordinate,
            "row_duration_error_s": max_duration,
            "objective_error_s": max_objective,
        },
    }


def _svg_document(lines: Sequence[str], *, width: int = 3200, height: int = 1800) -> str:
    line_height = 18
    if len(lines) * line_height + 70 > height:
        raise RuntimeError(f"Plain SVG preview has too many lines: {len(lines)}")
    body = [
        '<?xml version="1.0" encoding="UTF-8"?>',
        f'<svg xmlns="http://www.w3.org/2000/svg" width="{width}" height="{height}" viewBox="0 0 {width} {height}">',
        '<rect width="100%" height="100%" fill="white"/>',
    ]
    for index, line in enumerate(lines):
        size = 14 if index else 16
        fill = "#111111" if index else "#555555"
        body.append(
            f'<text x="40" y="{45 + index * line_height}" '
            f'font-family="Consolas, monospace" font-size="{size}" fill="{fill}" '
            f'xml:space="preserve">{html.escape(line)}</text>'
        )
    body.append("</svg>")
    return "\n".join(body) + "\n"


def _code_preview_lines(specs: Sequence[tuple[str, Path, int, int]]) -> list[str]:
    lines = ["Plain source preview; the paper uses searchable LaTeX, not this QA export."]
    for title, source, start, end in specs:
        plain_title = title.replace("\\quad", " ")
        lines.append(
            f"{plain_title} | {_relative(source)}:{start}-{end} | "
            f"sha256:{_sha256(source)[:16]}"
        )
        selected = source.read_text(encoding="utf-8").splitlines()[start - 1 : end]
        lines.extend(f"{number:4d}  {line}" for number, line in enumerate(selected, start))
        lines.append("")
    return lines


def _workbook_preview_lines(paths: Sequence[Path]) -> list[str]:
    lines = ["Plain workbook-value preview; the paper uses native three-line tables."]
    for path in paths:
        _sheet, rows = _read_active_rows(path)
        lines.append(f"{_relative(path)} | sha256:{_sha256(path)[:16]}")
        for row_index, row in enumerate(rows):
            if row_index == 0:
                values = [str(value).replace("烟幕干扰弹", "弹")[:18] for value in row]
            else:
                values = [
                    _format_number(value, 4) if isinstance(value, (int, float)) else str(value)
                    for value in row
                ]
            lines.append("\t".join(values))
        lines.append("")
    return lines


def _write_plain_svg_previews() -> list[dict[str, Any]]:
    result1 = ROOT / "outputs" / "result1.xlsx"
    result2 = ROOT / "outputs" / "result2.xlsx"
    result3 = ROOT / "outputs" / "result3.xlsx"
    specs: list[tuple[str, list[str], list[dict[str, Any]]]] = [
        (
            "appendix_code_geometry.svg",
            _code_preview_lines(CODE_FRAGMENT_SPECS["appendix_code_geometry.tex"]),
            [
                _code_snapshot(source, start, end)
                for _title, source, start, end in CODE_FRAGMENT_SPECS["appendix_code_geometry.tex"]
            ],
        ),
        (
            "appendix_code_optimizer.svg",
            _code_preview_lines(
                CODE_FRAGMENT_SPECS["appendix_code_evaluation.tex"]
                + CODE_FRAGMENT_SPECS["appendix_code_q5.tex"]
            ),
            [
                _code_snapshot(source, start, end)
                for _title, source, start, end in (
                    CODE_FRAGMENT_SPECS["appendix_code_evaluation.tex"]
                    + CODE_FRAGMENT_SPECS["appendix_code_q5.tex"]
                )
            ],
        ),
        (
            "appendix_data_result12.svg",
            _workbook_preview_lines((result1, result2)),
            [_workbook_snapshot(result1), _workbook_snapshot(result2)],
        ),
        (
            "appendix_data_result3.svg",
            _workbook_preview_lines((result3,)),
            [_workbook_snapshot(result3)],
        ),
    ]
    records: list[dict[str, Any]] = []
    for filename, lines, sources in specs:
        output = OUTPUT_DIR / filename
        output.write_text(_svg_document(lines), encoding="utf-8")
        records.append(
            {
                "path": _relative(output),
                "sha256": _sha256(output),
                "evidence_type": "plain_vector_qa_preview_not_embedded_in_paper",
                "width_px": 3200,
                "height_px": 1800,
                "sources": sources,
            }
        )
    return records


def _remove_legacy_posters() -> list[str]:
    removed: list[str] = []
    for filename in LEGACY_POSTERS:
        path = OUTPUT_DIR / filename
        if path.is_file():
            path.unlink()
            removed.append(_relative(path))
    return removed


def _write_manifest(
    *,
    artifacts: list[dict[str, Any]],
    fragments: list[dict[str, Any]],
    workbook_metrics: dict[str, Any],
    validation_metrics: dict[str, Any],
) -> dict[str, Any]:
    q35_path = ROOT / "validation" / "q3_q5_independent.json"
    q35 = _load_json(q35_path)
    workbooks = {
        name: ROOT / "outputs" / name
        for name in ("result1.xlsx", "result2.xlsx", "result3.xlsx")
    }
    expected_hashes = q35["output_workbook_sha256"]
    workbook_hash_match = {
        name: _sha256(path) == str(expected_hashes[name]).lower()
        for name, path in workbooks.items()
    }
    expected_rows = {"result1.xlsx": 3, "result2.xlsx": 3, "result3.xlsx": 15}
    row_counts = workbook_metrics["row_counts"]
    duration_sums = workbook_metrics["duration_sums"]
    exact_objectives = {
        question: float(q35["results"][question]["exact_centerline_objective"])
        for question in ("Q3", "Q4", "Q5")
    }
    claim_checks = {
        "workbook_hashes_match_validation": all(workbook_hash_match.values()),
        "active_row_counts_match_submission_scope": row_counts == expected_rows,
        "q4_nonoverlap_sum_matches_union": abs(duration_sums["result2.xlsx"] - exact_objectives["Q4"]) <= 1.0e-8,
        "q3_union_not_greater_than_row_sum": exact_objectives["Q3"] <= duration_sums["result1.xlsx"] + 1.0e-8,
        "q5_union_not_greater_than_row_sum": exact_objectives["Q5"] <= duration_sums["result3.xlsx"] + 1.0e-8,
        "roundtrip_errors_within_tolerance": max(validation_metrics["roundtrip_maxima"].values()) <= 1.0e-8,
    }
    if not all(claim_checks.values()):
        failed = [name for name, passed in claim_checks.items() if not passed]
        raise RuntimeError("Appendix claim checks failed: " + ", ".join(failed))

    requirements = ROOT / "requirements-formal.txt"
    renderer = Path(__file__).resolve()
    manifest = {
        "schema_version": 2,
        "status": "reproducible",
        "appendix_policy": {
            "paper_format": "searchable native LaTeX source and booktabs tables",
            "screenshots_embedded_in_paper": 0,
            "qa_previews": "plain SVG, not embedded",
        },
        "generator": {
            "path": _relative(renderer),
            "sha256": _sha256(renderer),
            "command": "python -B src/render_appendix_evidence.py",
            "python": sys.version.split()[0],
            "libraries": {"openpyxl": importlib.metadata.version("openpyxl")},
            "requirements_path": _relative(requirements),
            "requirements_sha256": _sha256(requirements),
        },
        "source_validation": {
            "path": _relative(q35_path),
            "sha256": _sha256(q35_path),
            "workbook_hash_match": workbook_hash_match,
            "active_data_rows": row_counts,
            "duration_row_sums_s": duration_sums,
            "exact_union_objectives_s": exact_objectives,
            "tolerance": 1.0e-8,
            "claim_checks": claim_checks,
        },
        "artifacts": artifacts,
        "latex_fragments": fragments,
        "validation_metrics": validation_metrics,
    }
    MANIFEST_PATH.write_text(
        json.dumps(manifest, ensure_ascii=False, indent=2, sort_keys=True) + "\n",
        encoding="utf-8",
    )
    return manifest


def main() -> None:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    file_index = _write_file_index()
    code_fragments = _write_code_fragments()
    table_fragments, workbook_metrics = _write_table_rows()
    validation_fragments, validation_metrics = _write_validation_rows()
    artifacts = _write_plain_svg_previews()
    removed = _remove_legacy_posters()
    fragments = [file_index] + code_fragments + table_fragments + validation_fragments
    manifest = _write_manifest(
        artifacts=artifacts,
        fragments=fragments,
        workbook_metrics=workbook_metrics,
        validation_metrics=validation_metrics,
    )

    for record in fragments:
        print(f"{record['path']}  sha256:{record['sha256']}")
    for record in artifacts:
        print(f"{record['path']}  sha256:{record['sha256']}")
    for path in removed:
        print(f"removed legacy poster: {path}")
    print(f"{_relative(MANIFEST_PATH)}  sha256:{_sha256(MANIFEST_PATH)}")
    print(
        "appendix policy: "
        f"{manifest['appendix_policy']['screenshots_embedded_in_paper']} embedded screenshots"
    )


if __name__ == "__main__":
    main()
