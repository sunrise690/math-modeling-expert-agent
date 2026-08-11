from __future__ import annotations

import argparse
import json
import re
from pathlib import Path
from typing import Any


def _safe_name(value: str, default: str = "report") -> str:
    name = re.sub(r"[^A-Za-z0-9._-]+", "-", value.strip()).strip("-.")
    return name[:80] or default


def _set_font(style: Any, western: str, east_asia: str, size: float | None = None) -> None:
    from docx.oxml.ns import qn
    from docx.shared import Pt

    style.font.name = western
    style._element.rPr.rFonts.set(qn("w:eastAsia"), east_asia)
    if size:
        style.font.size = Pt(size)


def _add_markdown(document: Any, markdown: str) -> None:
    lines = markdown.splitlines()
    index = 0
    in_code = False
    code_lines: list[str] = []
    while index < len(lines):
        line = lines[index]
        if line.strip().startswith("```"):
            if in_code:
                paragraph = document.add_paragraph("\n".join(code_lines))
                paragraph.style = document.styles["No Spacing"]
                code_lines = []
            in_code = not in_code
            index += 1
            continue
        if in_code:
            code_lines.append(line)
            index += 1
            continue
        heading = re.match(r"^(#{1,4})\s+(.+)$", line)
        if heading:
            document.add_heading(heading.group(2).strip(), level=min(len(heading.group(1)), 4))
        elif re.match(r"^[-*]\s+", line):
            document.add_paragraph(re.sub(r"^[-*]\s+", "", line), style="List Bullet")
        elif re.match(r"^\d+[.]\s+", line):
            document.add_paragraph(re.sub(r"^\d+[.]\s+", "", line), style="List Number")
        elif line.strip().startswith("|") and index + 1 < len(lines) and re.match(r"^\s*\|?\s*:?-+", lines[index + 1]):
            headers = [cell.strip() for cell in line.strip().strip("|").split("|")]
            index += 2
            rows: list[list[str]] = []
            while index < len(lines) and lines[index].strip().startswith("|"):
                rows.append([cell.strip() for cell in lines[index].strip().strip("|").split("|")])
                index += 1
            table = document.add_table(rows=1, cols=len(headers))
            table.style = "Table Grid"
            for cell, value in zip(table.rows[0].cells, headers):
                cell.text = value
            for values in rows:
                cells = table.add_row().cells
                for cell, value in zip(cells, values):
                    cell.text = value
            continue
        elif line.strip():
            document.add_paragraph(line.strip())
        index += 1


def export_report(spec: dict[str, Any], output_dir: Path) -> dict[str, Any]:
    markdown = str(spec.get("markdown", "")).strip()
    if not markdown:
        raise ValueError("markdown 不能为空")
    if len(markdown) > 500_000:
        raise ValueError("markdown 过长")

    output_dir.mkdir(parents=True, exist_ok=True)
    filename = _safe_name(str(spec.get("filename", "report")))
    formats = spec.get("formats", ["md", "docx"])
    if not isinstance(formats, list):
        formats = ["md", "docx"]
    formats = list(dict.fromkeys(str(item).lower() for item in formats if str(item).lower() in {"md", "docx"}))
    artifacts: list[str] = []

    if "md" in formats:
        markdown_path = output_dir / f"{filename}.md"
        markdown_path.write_text(markdown + "\n", encoding="utf-8")
        artifacts.append(str(markdown_path.resolve()))

    if "docx" in formats:
        from docx import Document
        from docx.shared import Cm

        document = Document()
        section = document.sections[0]
        section.top_margin = Cm(2.54)
        section.bottom_margin = Cm(2.54)
        section.left_margin = Cm(2.7)
        section.right_margin = Cm(2.7)
        _set_font(document.styles["Normal"], "Times New Roman", "宋体", 10.5)
        for level in range(1, 5):
            _set_font(document.styles[f"Heading {level}"], "Arial", "黑体")
        title = str(spec.get("title", "")).strip()
        if title and not markdown.lstrip().startswith("#"):
            document.add_heading(title, 0)
        _add_markdown(document, markdown)
        docx_path = output_dir / f"{filename}.docx"
        document.save(docx_path)
        artifacts.append(str(docx_path.resolve()))

    return {"ok": True, "formats": formats, "artifacts": artifacts}


def main() -> None:
    parser = argparse.ArgumentParser(description="Export Markdown to editable report files.")
    parser.add_argument("--spec", type=Path, required=True)
    parser.add_argument("--output-dir", type=Path, required=True)
    args = parser.parse_args()
    spec = json.loads(args.spec.read_text(encoding="utf-8-sig"))
    if not isinstance(spec, dict):
        raise ValueError("spec 必须是 JSON 对象")
    print(json.dumps(export_report(spec, args.output_dir), ensure_ascii=False))


if __name__ == "__main__":
    main()
