from __future__ import annotations

import hashlib
import json
import re
import sqlite3
import threading
import uuid
from contextlib import contextmanager
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Iterator


SUPPORTED_MATERIAL_EXTENSIONS = {
    ".bib",
    ".bst",
    ".csv",
    ".cls",
    ".docx",
    ".json",
    ".m",
    ".md",
    ".pdf",
    ".py",
    ".tex",
    ".tsv",
    ".txt",
    ".xlsx",
}
TEXT_MATERIAL_EXTENSIONS = {".bib", ".bst", ".cls", ".json", ".m", ".md", ".py", ".tex", ".txt"}
SENSITIVE_SPREADSHEET_NAME = re.compile(r"参赛信息|参赛队|报名|名单|通讯录|汇总", re.IGNORECASE)
SENSITIVE_COLUMN_NAME = re.compile(
    r"姓名|学号|电话|手机|邮箱|email|身份证|证件号|住址|地址|微信|qq|指导教师|队员",
    re.IGNORECASE,
)
MAX_MATERIAL_BYTES = 100 * 1024 * 1024
MAX_DOCUMENT_CHARS = 2_000_000
CHUNK_CHARS = 1_800
CHUNK_OVERLAP = 180


class KnowledgeError(RuntimeError):
    pass


class MaterialExcluded(KnowledgeError):
    pass


def _utc_now() -> str:
    return datetime.now(timezone.utc).isoformat()


class KnowledgeBase:
    """Incremental, local-only full-text index for modeling reference material."""

    def __init__(self, database: Path, roots: list[Path]) -> None:
        self.database = database
        self.database.parent.mkdir(parents=True, exist_ok=True)
        self.roots = self._normalize_roots(roots)
        self._index_lock = threading.Lock()
        self._job_lock = threading.Lock()
        self._job_thread: threading.Thread | None = None
        self._job: dict[str, Any] = {"status": "idle"}
        self._initialize()

    @staticmethod
    def _normalize_roots(roots: list[Path]) -> list[Path]:
        normalized: list[Path] = []
        seen: set[str] = set()
        for raw_root in roots:
            try:
                root = raw_root.expanduser().resolve()
            except OSError:
                continue
            key = str(root).casefold()
            if key in seen:
                continue
            seen.add(key)
            normalized.append(root)
        return normalized

    @contextmanager
    def _connect(self) -> Iterator[sqlite3.Connection]:
        connection = sqlite3.connect(self.database, timeout=30)
        connection.row_factory = sqlite3.Row
        try:
            yield connection
            connection.commit()
        except Exception:
            connection.rollback()
            raise
        finally:
            connection.close()

    def _initialize(self) -> None:
        with self._connect() as connection:
            connection.execute("PRAGMA journal_mode=WAL")
            connection.executescript(
                """
                CREATE TABLE IF NOT EXISTS knowledge_meta (
                    key TEXT PRIMARY KEY,
                    value TEXT NOT NULL
                );
                CREATE TABLE IF NOT EXISTS documents (
                    id TEXT PRIMARY KEY,
                    root_path TEXT NOT NULL,
                    absolute_path TEXT NOT NULL UNIQUE,
                    relative_path TEXT NOT NULL,
                    name TEXT NOT NULL,
                    extension TEXT NOT NULL,
                    size INTEGER NOT NULL,
                    mtime_ns INTEGER NOT NULL,
                    fingerprint TEXT NOT NULL,
                    status TEXT NOT NULL,
                    error TEXT NOT NULL DEFAULT '',
                    page_count INTEGER NOT NULL DEFAULT 0,
                    char_count INTEGER NOT NULL DEFAULT 0,
                    chunk_count INTEGER NOT NULL DEFAULT 0,
                    content_hash TEXT NOT NULL DEFAULT '',
                    duplicate_of TEXT NOT NULL DEFAULT '',
                    indexed_at TEXT NOT NULL
                );
                CREATE INDEX IF NOT EXISTS idx_documents_status ON documents(status);
                CREATE INDEX IF NOT EXISTS idx_documents_extension ON documents(extension);
                """
            )
            columns = {str(row[1]) for row in connection.execute("PRAGMA table_info(documents)").fetchall()}
            if "content_hash" not in columns:
                connection.execute("ALTER TABLE documents ADD COLUMN content_hash TEXT NOT NULL DEFAULT ''")
            if "duplicate_of" not in columns:
                connection.execute("ALTER TABLE documents ADD COLUMN duplicate_of TEXT NOT NULL DEFAULT ''")
            try:
                connection.execute(
                    """
                    CREATE VIRTUAL TABLE IF NOT EXISTS material_chunks USING fts5(
                        document_id UNINDEXED,
                        ordinal UNINDEXED,
                        location UNINDEXED,
                        name,
                        relative_path,
                        text,
                        tokenize='trigram'
                    )
                    """
                )
                tokenizer = "trigram"
            except sqlite3.OperationalError:
                connection.execute(
                    """
                    CREATE VIRTUAL TABLE IF NOT EXISTS material_chunks USING fts5(
                        document_id UNINDEXED,
                        ordinal UNINDEXED,
                        location UNINDEXED,
                        name,
                        relative_path,
                        text,
                        tokenize='unicode61'
                    )
                    """
                )
                tokenizer = "unicode61"
            connection.execute(
                "INSERT OR REPLACE INTO knowledge_meta(key, value) VALUES ('tokenizer', ?)",
                (tokenizer,),
            )

    def status(self) -> dict[str, Any]:
        with self._connect() as connection:
            rows = connection.execute(
                "SELECT status, COUNT(*) AS count, COALESCE(SUM(chunk_count), 0) AS chunks, "
                "COALESCE(SUM(char_count), 0) AS chars FROM documents GROUP BY status"
            ).fetchall()
            extensions = connection.execute(
                "SELECT extension, COUNT(*) AS count FROM documents GROUP BY extension ORDER BY count DESC, extension"
            ).fetchall()
            meta = {
                str(row["key"]): str(row["value"])
                for row in connection.execute("SELECT key, value FROM knowledge_meta").fetchall()
            }
        by_status = {str(row["status"]): int(row["count"]) for row in rows}
        return {
            "ready": by_status.get("indexed", 0) > 0,
            "roots": [{"path": str(root), "exists": root.is_dir()} for root in self.roots],
            "documents": sum(by_status.values()),
            "indexedDocuments": by_status.get("indexed", 0),
            "emptyDocuments": by_status.get("empty", 0),
            "failedDocuments": by_status.get("error", 0),
            "excludedDocuments": by_status.get("excluded", 0),
            "duplicateDocuments": by_status.get("duplicate", 0),
            "chunks": sum(int(row["chunks"]) for row in rows),
            "characters": sum(int(row["chars"]) for row in rows),
            "extensions": {str(row["extension"]): int(row["count"]) for row in extensions},
            "lastIndexedAt": meta.get("last_indexed_at", ""),
            "tokenizer": meta.get("tokenizer", ""),
            "supportedExtensions": sorted(SUPPORTED_MATERIAL_EXTENSIONS),
            "indexJob": self.index_job_status(),
        }

    def start_reindex(self) -> dict[str, Any]:
        with self._job_lock:
            if self._job_thread is not None and self._job_thread.is_alive():
                return dict(self._job)
            job_id = uuid.uuid4().hex
            self._job = {
                "id": job_id,
                "status": "running",
                "startedAt": _utc_now(),
                "finishedAt": "",
                "error": "",
            }
            self._job_thread = threading.Thread(
                target=self._run_reindex_job,
                args=(job_id,),
                daemon=True,
                name="knowledge-indexer",
            )
            self._job_thread.start()
            return dict(self._job)

    def _run_reindex_job(self, job_id: str) -> None:
        try:
            result = self.reindex()
        except Exception as error:
            with self._job_lock:
                if self._job.get("id") == job_id:
                    self._job.update(
                        {
                            "status": "failed",
                            "finishedAt": _utc_now(),
                            "error": str(error).strip() or type(error).__name__,
                        }
                    )
        else:
            with self._job_lock:
                if self._job.get("id") == job_id:
                    self._job.update(
                        {
                            "status": "completed",
                            "finishedAt": _utc_now(),
                            "error": "",
                            "result": {key: value for key, value in result.items() if key != "status"},
                        }
                    )

    def index_job_status(self) -> dict[str, Any]:
        with self._job_lock:
            return dict(self._job)

    def reindex(self) -> dict[str, Any]:
        if not self._index_lock.acquire(blocking=False):
            raise KnowledgeError("资料库正在建立索引")
        try:
            started_at = _utc_now()
            discovered = self._discover_files()
            discovered_paths = {str(path.resolve()).casefold() for _, path in discovered}
            with self._connect() as connection:
                existing = {
                    str(row["absolute_path"]).casefold(): row
                    for row in connection.execute("SELECT * FROM documents").fetchall()
                }

            summary = {
                "discovered": len(discovered),
                "indexed": 0,
                "unchanged": 0,
                "empty": 0,
                "failed": 0,
                "excluded": 0,
                "duplicates": 0,
                "removed": 0,
            }
            existing_by_id = {str(row["id"]): row for row in existing.values()}
            hash_owners: dict[str, str] = {}
            for row in existing.values():
                if (
                    str(row["status"]) == "indexed"
                    and str(row["content_hash"])
                    and str(row["absolute_path"]).casefold() in discovered_paths
                ):
                    hash_owners.setdefault(str(row["content_hash"]), str(row["id"]))
            for root, path in discovered:
                absolute = path.resolve()
                key = str(absolute).casefold()
                stat = absolute.stat()
                fingerprint = f"{stat.st_size}:{stat.st_mtime_ns}"
                current = existing.get(key)
                duplicate_owner_available = True
                if current is not None and str(current["status"]) == "duplicate":
                    owner = existing_by_id.get(str(current["duplicate_of"]))
                    duplicate_owner_available = bool(
                        owner is not None and str(owner["absolute_path"]).casefold() in discovered_paths
                    )
                if (
                    current is not None
                    and str(current["fingerprint"]) == fingerprint
                    and str(current["status"]) != "error"
                    and duplicate_owner_available
                ):
                    summary["unchanged"] += 1
                    continue
                content_hash = self._file_hash(absolute)
                document_id = self._document_id(absolute)
                duplicate_of = hash_owners.get(content_hash, "")
                if duplicate_of and duplicate_of != document_id:
                    owner = existing_by_id.get(duplicate_of)
                    owner_name = str(owner["relative_path"]) if owner is not None else duplicate_of
                    self._replace_document(
                        root,
                        absolute,
                        fingerprint,
                        "duplicate",
                        f"与 {owner_name} 内容重复，已避免重复索引",
                        0,
                        0,
                        [],
                        content_hash,
                        duplicate_of,
                    )
                    summary["duplicates"] += 1
                    continue
                try:
                    extracted, page_count = self._extract_document(absolute)
                    chunks = self._make_chunks(extracted)
                    status = "indexed" if chunks else "empty"
                    error = "" if chunks else "未提取到文本，文件可能是扫描件或仅含图片"
                    char_count = sum(len(text) for _, text in extracted)
                    self._replace_document(
                        root,
                        absolute,
                        fingerprint,
                        status,
                        error,
                        page_count,
                        char_count,
                        chunks,
                        content_hash,
                        "",
                    )
                    if chunks:
                        hash_owners.setdefault(content_hash, document_id)
                    summary["indexed" if chunks else "empty"] += 1
                except MaterialExcluded as error:
                    self._replace_document(
                        root,
                        absolute,
                        fingerprint,
                        "excluded",
                        str(error).strip()[:1000],
                        0,
                        0,
                        [],
                        content_hash,
                        "",
                    )
                    summary["excluded"] += 1
                except Exception as error:
                    self._replace_document(
                        root,
                        absolute,
                        fingerprint,
                        "error",
                        str(error).strip()[:1000] or type(error).__name__,
                        0,
                        0,
                        [],
                        content_hash,
                        "",
                    )
                    summary["failed"] += 1

            stale = [row for key, row in existing.items() if key not in discovered_paths]
            if stale:
                with self._connect() as connection:
                    for row in stale:
                        connection.execute("DELETE FROM material_chunks WHERE document_id = ?", (row["id"],))
                        connection.execute("DELETE FROM documents WHERE id = ?", (row["id"],))
                summary["removed"] = len(stale)

            finished_at = _utc_now()
            with self._connect() as connection:
                connection.execute(
                    "INSERT OR REPLACE INTO knowledge_meta(key, value) VALUES ('last_indexed_at', ?)",
                    (finished_at,),
                )
                connection.execute(
                    "INSERT OR REPLACE INTO knowledge_meta(key, value) VALUES ('last_index_summary', ?)",
                    (json.dumps(summary, ensure_ascii=False),),
                )
            return {"startedAt": started_at, "finishedAt": finished_at, **summary, "status": self.status()}
        finally:
            self._index_lock.release()

    def _discover_files(self) -> list[tuple[Path, Path]]:
        discovered: list[tuple[Path, Path]] = []
        for root in self.roots:
            if not root.is_dir():
                continue
            for path in root.rglob("*"):
                if not path.is_file() or path.suffix.lower() not in SUPPORTED_MATERIAL_EXTENSIONS:
                    continue
                try:
                    resolved = path.resolve()
                    resolved.relative_to(root)
                    size = resolved.stat().st_size
                except OSError:
                    continue
                except ValueError:
                    continue
                if 0 < size <= MAX_MATERIAL_BYTES:
                    discovered.append((root, resolved))
        discovered.sort(key=lambda item: str(item[1]).casefold())
        return discovered

    @staticmethod
    def _document_id(path: Path) -> str:
        return hashlib.sha256(str(path).casefold().encode("utf-8")).hexdigest()[:32]

    @staticmethod
    def _file_hash(path: Path) -> str:
        digest = hashlib.sha256()
        with path.open("rb") as handle:
            for block in iter(lambda: handle.read(1024 * 1024), b""):
                digest.update(block)
        return digest.hexdigest()

    def _replace_document(
        self,
        root: Path,
        path: Path,
        fingerprint: str,
        status: str,
        error: str,
        page_count: int,
        char_count: int,
        chunks: list[tuple[int, str, str]],
        content_hash: str,
        duplicate_of: str,
    ) -> None:
        document_id = self._document_id(path)
        relative_path = str(path.relative_to(root))
        stat = path.stat()
        with self._connect() as connection:
            connection.execute("DELETE FROM material_chunks WHERE document_id = ?", (document_id,))
            connection.execute(
                """
                INSERT INTO documents(
                    id, root_path, absolute_path, relative_path, name, extension, size,
                    mtime_ns, fingerprint, status, error, page_count, char_count,
                    chunk_count, content_hash, duplicate_of, indexed_at
                ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                ON CONFLICT(id) DO UPDATE SET
                    root_path=excluded.root_path,
                    absolute_path=excluded.absolute_path,
                    relative_path=excluded.relative_path,
                    name=excluded.name,
                    extension=excluded.extension,
                    size=excluded.size,
                    mtime_ns=excluded.mtime_ns,
                    fingerprint=excluded.fingerprint,
                    status=excluded.status,
                    error=excluded.error,
                    page_count=excluded.page_count,
                    char_count=excluded.char_count,
                    chunk_count=excluded.chunk_count,
                    content_hash=excluded.content_hash,
                    duplicate_of=excluded.duplicate_of,
                    indexed_at=excluded.indexed_at
                """,
                (
                    document_id,
                    str(root),
                    str(path),
                    relative_path,
                    path.name,
                    path.suffix.lower(),
                    stat.st_size,
                    stat.st_mtime_ns,
                    fingerprint,
                    status,
                    error,
                    page_count,
                    char_count,
                    len(chunks),
                    content_hash,
                    duplicate_of,
                    _utc_now(),
                ),
            )
            connection.executemany(
                """
                INSERT INTO material_chunks(document_id, ordinal, location, name, relative_path, text)
                VALUES (?, ?, ?, ?, ?, ?)
                """,
                [
                    (document_id, ordinal, location, path.name, relative_path, text)
                    for ordinal, location, text in chunks
                ],
            )

    def _extract_document(self, path: Path) -> tuple[list[tuple[str, str]], int]:
        extension = path.suffix.lower()
        if extension == ".pdf":
            return self._extract_pdf(path)
        if extension == ".docx":
            return self._extract_docx(path), 0
        if extension == ".xlsx":
            if SENSITIVE_SPREADSHEET_NAME.search(path.stem):
                raise MaterialExcluded("文件名表明可能包含参赛者个人信息，默认不建立全文索引")
            return self._extract_xlsx(path), 0
        if extension in {".csv", ".tsv"}:
            return self._extract_delimited(path), 0
        if extension in TEXT_MATERIAL_EXTENSIONS:
            return [("正文", self._read_text(path))], 0
        raise KnowledgeError(f"不支持的资料格式：{extension}")

    @staticmethod
    def _extract_pdf(path: Path) -> tuple[list[tuple[str, str]], int]:
        from pypdf import PdfReader

        reader = PdfReader(str(path), strict=False)
        if reader.is_encrypted:
            try:
                reader.decrypt("")
            except Exception as error:
                raise KnowledgeError("PDF 已加密，无法提取文本") from error
        extracted: list[tuple[str, str]] = []
        for page_number, page in enumerate(reader.pages, start=1):
            text = page.extract_text() or ""
            if text.strip():
                extracted.append((f"第 {page_number} 页", text))
        return extracted, len(reader.pages)

    @staticmethod
    def _extract_docx(path: Path) -> list[tuple[str, str]]:
        from docx import Document

        document = Document(path)
        parts = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
        for table_number, table in enumerate(document.tables, start=1):
            rows = []
            for row in table.rows:
                rows.append(" | ".join(cell.text.strip() for cell in row.cells))
            if rows:
                parts.append(f"[表 {table_number}]\n" + "\n".join(rows))
        return [("正文", "\n".join(parts))]

    @staticmethod
    def _frame_to_text(frame: Any, max_rows: int = 20_000) -> str:
        if len(frame) > max_rows:
            frame = frame.head(max_rows)
        frame = frame.fillna("")
        return frame.to_csv(index=False, lineterminator="\n")

    def _extract_xlsx(self, path: Path) -> list[tuple[str, str]]:
        import pandas as pd

        extracted: list[tuple[str, str]] = []
        with pd.ExcelFile(path) as workbook:
            for sheet_name in workbook.sheet_names[:50]:
                frame = pd.read_excel(workbook, sheet_name=sheet_name)
                sensitive_columns = [column for column in frame.columns if SENSITIVE_COLUMN_NAME.search(str(column))]
                if sensitive_columns:
                    frame = frame.drop(columns=sensitive_columns)
                if not frame.empty or len(frame.columns):
                    extracted.append((f"工作表：{sheet_name}", self._frame_to_text(frame)))
        return extracted

    def _extract_delimited(self, path: Path) -> list[tuple[str, str]]:
        import pandas as pd

        separator = "\t" if path.suffix.lower() == ".tsv" else None
        last_error: Exception | None = None
        for encoding in ("utf-8-sig", "utf-8", "gb18030"):
            try:
                frame = pd.read_csv(path, sep=separator, engine="python", encoding=encoding)
                return [("数据表", self._frame_to_text(frame))]
            except UnicodeDecodeError as error:
                last_error = error
        raise KnowledgeError("无法识别文本表格编码") from last_error

    @staticmethod
    def _read_text(path: Path) -> str:
        last_error: Exception | None = None
        for encoding in ("utf-8-sig", "utf-8", "gb18030"):
            try:
                return path.read_text(encoding=encoding)
            except UnicodeDecodeError as error:
                last_error = error
        raise KnowledgeError("无法识别文本文件编码") from last_error

    def _make_chunks(self, extracted: list[tuple[str, str]]) -> list[tuple[int, str, str]]:
        chunks: list[tuple[int, str, str]] = []
        total_chars = 0
        ordinal = 0
        for location, raw_text in extracted:
            text = self._normalize_text(raw_text)
            if not text:
                continue
            remaining = MAX_DOCUMENT_CHARS - total_chars
            if remaining <= 0:
                break
            text = text[:remaining]
            total_chars += len(text)
            for chunk in self._chunk_text(text):
                chunks.append((ordinal, location, chunk))
                ordinal += 1
        return chunks

    @staticmethod
    def _normalize_text(text: str) -> str:
        text = str(text).replace("\x00", " ").replace("\r\n", "\n").replace("\r", "\n")
        # Some PDFs expose malformed UTF-16 glyphs as isolated surrogates. SQLite's
        # UTF-8 encoder rejects those code points, so preserve the surrounding text
        # and mark only the invalid glyph instead of failing the whole document.
        text = re.sub(r"[\ud800-\udfff]", "\ufffd", text)
        text = re.sub(r"[\t\f\v ]+", " ", text)
        text = re.sub(r"\n{3,}", "\n\n", text)
        return text.strip()

    @staticmethod
    def _chunk_text(text: str) -> list[str]:
        if len(text) <= CHUNK_CHARS:
            return [text]
        chunks: list[str] = []
        start = 0
        while start < len(text):
            end = min(len(text), start + CHUNK_CHARS)
            if end < len(text):
                floor = start + CHUNK_CHARS // 2
                candidates = [text.rfind(mark, floor, end) for mark in ("\n", "。", "！", "？", ". ", "; ")]
                boundary = max(candidates)
                if boundary > floor:
                    end = boundary + 1
            chunk = text[start:end].strip()
            if chunk:
                chunks.append(chunk)
            if end >= len(text):
                break
            start = max(start + 1, end - CHUNK_OVERLAP)
        return chunks

    def search(
        self,
        query: str,
        limit: int = 5,
        extensions: list[str] | None = None,
    ) -> dict[str, Any]:
        normalized_query = " ".join(str(query).strip().split())
        if not normalized_query:
            raise KnowledgeError("检索词不能为空")
        if len(normalized_query) > 200:
            raise KnowledgeError("检索词不能超过 200 个字符")
        limit = max(1, min(10, int(limit)))
        allowed_extensions = self._normalize_extensions(extensions)
        terms = self._query_terms(normalized_query)
        candidates = self._fts_candidates(terms, limit * 12, allowed_extensions)
        if not candidates:
            candidates = self._like_candidates(normalized_query, limit * 12, allowed_extensions)

        ranked = []
        query_lower = normalized_query.casefold()
        for row in candidates:
            text = str(row["text"])
            haystack = f"{row['name']}\n{row['relative_path']}\n{text}".casefold()
            matched_terms = [term for term in terms if term.casefold() in haystack]
            exact = query_lower in haystack
            filename_hits = sum(1 for term in terms if term.casefold() in str(row["name"]).casefold())
            lexical = len(matched_terms) / max(1, len(terms))
            rank_value = float(row["fts_rank"]) if row["fts_rank"] is not None else 0.0
            score = lexical + (2.0 if exact else 0.0) + filename_hits * 0.5 + max(0.0, -rank_value)
            ranked.append((score, row, matched_terms))
        ranked.sort(key=lambda item: (-item[0], str(item[1]["relative_path"]), int(item[1]["ordinal"])))

        results = []
        per_document: dict[str, int] = {}
        for score, row, matched_terms in ranked:
            document_id = str(row["document_id"])
            if per_document.get(document_id, 0) >= 2:
                continue
            per_document[document_id] = per_document.get(document_id, 0) + 1
            results.append(
                {
                    "documentId": document_id,
                    "chunkId": int(row["chunk_id"]),
                    "name": str(row["name"]),
                    "relativePath": str(row["relative_path"]),
                    "collection": Path(str(row["root_path"])).name or "local",
                    "extension": str(row["extension"]),
                    "contentHash": str(row["content_hash"]),
                    "location": str(row["location"]),
                    "citation": f"{row['name']}（{row['location']}）",
                    "excerpt": self._excerpt(str(row["text"]), matched_terms or terms),
                    "score": round(score, 6),
                }
            )
            if len(results) >= limit:
                break
        return {"query": normalized_query, "count": len(results), "results": results}

    @staticmethod
    def _normalize_extensions(extensions: list[str] | None) -> list[str]:
        if extensions is None:
            return []
        if not isinstance(extensions, list):
            raise KnowledgeError("extensions 必须是扩展名数组")
        normalized = []
        for item in extensions:
            extension = str(item).strip().lower()
            if extension and not extension.startswith("."):
                extension = f".{extension}"
            if extension in SUPPORTED_MATERIAL_EXTENSIONS and extension not in normalized:
                normalized.append(extension)
        return normalized

    @staticmethod
    def _query_terms(query: str) -> list[str]:
        raw_terms = re.findall(r"[A-Za-z0-9_+.-]+|[\u3400-\u9fff]+", query)
        terms: list[str] = []
        for raw in raw_terms:
            term = raw.strip().casefold()
            if not term:
                continue
            if re.fullmatch(r"[\u3400-\u9fff]+", term) and len(term) > 10:
                pieces = [term[index : index + 5] for index in range(0, len(term) - 2, 3)]
            else:
                pieces = [term]
            for piece in pieces:
                if piece not in terms:
                    terms.append(piece)
        return terms[:16] or [query.casefold()]

    def _fts_candidates(self, terms: list[str], limit: int, extensions: list[str]) -> list[sqlite3.Row]:
        searchable = [term for term in terms if len(term) >= 3]
        if not searchable:
            return []
        expression = " OR ".join(f'"{term.replace(chr(34), chr(34) * 2)}"' for term in searchable)
        extension_sql = ""
        parameters: list[Any] = [expression]
        if extensions:
            extension_sql = f" AND d.extension IN ({','.join('?' for _ in extensions)})"
            parameters.extend(extensions)
        parameters.append(limit)
        try:
            with self._connect() as connection:
                return connection.execute(
                    f"""
                    SELECT material_chunks.rowid AS chunk_id, material_chunks.document_id,
                           material_chunks.ordinal, material_chunks.location,
                           material_chunks.name, material_chunks.relative_path,
                           material_chunks.text, d.root_path, d.extension, d.content_hash,
                           bm25(material_chunks, 0.0, 0.0, 0.0, 4.0, 2.0, 1.0) AS fts_rank
                    FROM material_chunks
                    JOIN documents d ON d.id = material_chunks.document_id
                    WHERE material_chunks MATCH ? AND d.status = 'indexed'{extension_sql}
                    ORDER BY fts_rank
                    LIMIT ?
                    """,
                    parameters,
                ).fetchall()
        except sqlite3.OperationalError:
            return []

    def _like_candidates(self, query: str, limit: int, extensions: list[str]) -> list[sqlite3.Row]:
        extension_sql = ""
        parameters: list[Any] = [f"%{query}%", f"%{query}%", f"%{query}%"]
        if extensions:
            extension_sql = f" AND d.extension IN ({','.join('?' for _ in extensions)})"
            parameters.extend(extensions)
        parameters.append(limit)
        with self._connect() as connection:
            return connection.execute(
                f"""
                SELECT material_chunks.rowid AS chunk_id, material_chunks.document_id,
                       material_chunks.ordinal, material_chunks.location,
                       material_chunks.name, material_chunks.relative_path,
                       material_chunks.text, d.root_path, d.extension, d.content_hash, 0.0 AS fts_rank
                FROM material_chunks
                JOIN documents d ON d.id = material_chunks.document_id
                WHERE d.status = 'indexed'
                  AND (material_chunks.text LIKE ? OR material_chunks.name LIKE ? OR material_chunks.relative_path LIKE ?)
                  {extension_sql}
                LIMIT ?
                """,
                parameters,
            ).fetchall()

    @staticmethod
    def _excerpt(text: str, terms: list[str], max_chars: int = 650) -> str:
        lowered = text.casefold()
        positions = [lowered.find(term.casefold()) for term in terms if term and lowered.find(term.casefold()) >= 0]
        center = min(positions) if positions else 0
        start = max(0, center - max_chars // 3)
        end = min(len(text), start + max_chars)
        excerpt = text[start:end].strip()
        if start:
            excerpt = f"…{excerpt}"
        if end < len(text):
            excerpt = f"{excerpt}…"
        return excerpt

    def read(self, document_id: str, chunk_id: int | None = None, max_chars: int = 6_000) -> dict[str, Any]:
        if not re.fullmatch(r"[a-f0-9]{32}", str(document_id)):
            raise KnowledgeError("资料 ID 无效")
        max_chars = max(1_000, min(20_000, int(max_chars)))
        with self._connect() as connection:
            document = connection.execute("SELECT * FROM documents WHERE id = ?", (document_id,)).fetchone()
            if document is None:
                raise KnowledgeError("资料不存在")
            if chunk_id is None:
                rows = connection.execute(
                    """
                    SELECT rowid AS chunk_id, ordinal, location, text
                    FROM material_chunks WHERE document_id = ? ORDER BY CAST(ordinal AS INTEGER)
                    """,
                    (document_id,),
                ).fetchall()
            else:
                rows = connection.execute(
                    """
                    SELECT rowid AS chunk_id, ordinal, location, text
                    FROM material_chunks WHERE document_id = ? AND rowid = ?
                    """,
                    (document_id, int(chunk_id)),
                ).fetchall()
        if chunk_id is not None and not rows:
            raise KnowledgeError("资料片段不存在")
        selected = []
        used = 0
        truncated = False
        for row in rows:
            text = str(row["text"])
            remaining = max_chars - used
            if remaining <= 0:
                truncated = True
                break
            selected_text = text[:remaining]
            selected.append(
                {
                    "chunkId": int(row["chunk_id"]),
                    "ordinal": int(row["ordinal"]),
                    "location": str(row["location"]),
                    "text": selected_text,
                }
            )
            used += len(selected_text)
            if len(selected_text) < len(text):
                truncated = True
                break
        return {
            "document": {
                "id": str(document["id"]),
                "collection": Path(str(document["root_path"])).name or "local",
                "name": str(document["name"]),
                "relativePath": str(document["relative_path"]),
                "extension": str(document["extension"]),
                "contentHash": str(document["content_hash"]),
                "pages": int(document["page_count"]),
                "characters": int(document["char_count"]),
                "chunks": int(document["chunk_count"]),
                "status": str(document["status"]),
                "error": str(document["error"]),
            },
            "content": selected,
            "truncated": truncated,
        }
